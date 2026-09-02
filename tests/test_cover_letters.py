import json

import pytest
from sqlmodel import Session, select
from tests.fakes import scripted_api_fn

from findmyjob.docx.render import render_cover_letter
from findmyjob.llm.client import LlmClient
from findmyjob.llm.cover_letter import CoverLetterResult
from findmyjob.models.application import Application, CoverLetter
from findmyjob.models.enums import (
    ApplicationStatus,
    ContractType,
    CoverLetterLanguageMode,
)
from findmyjob.models.job import Job, JobAnalysis
from findmyjob.services.analyze import ANALYZER_VERSION
from findmyjob.services.cover_letters import (
    _resolve_language,
    generate,
    update,
)
from findmyjob.services.settings import update_app_settings

pytestmark = [pytest.mark.usefixtures("seeded_session"), pytest.mark.slow]

LETTER = {
    "language": "de",
    "recipient": {"company": "Acme GmbH"},
    "subject": "Bewerbung als Werkstudent Data Science",
    "salutation": "Sehr geehrte Damen und Herren,",
    "paragraphs": [
        "Ihre Stellenanzeige für die Arbeit an Datenpipelines hat mich sofort angesprochen.",
        "In drei Jahren mit Python habe ich mehrere ETL-Projekte umgesetzt.",
        "Ich bin eingeschrieben und kann 20 Stunden pro Woche einbringen.",
    ],
    "closing": "Mit freundlichen Grüßen",
    "claims_used": [{"claim": "3 years Python", "evidence_from_profile": "CV: Python since 2023"}],
}


def _client() -> LlmClient:
    return LlmClient(api_fn=scripted_api_fn(json.dumps(LETTER)))


def _job_with_analysis(db: Session, *, english_only: bool = False) -> int:
    job = Job(
        source_key="s",
        source_job_id="1",
        url="https://x.test/1",
        title="Werkstudent Data Science",
        normalized_title="werkstudent data science",
        company_name_raw="Acme GmbH",
        jd_text="Werkstudent role, Python and SQL, 20h/week, enrolled students only.",
    )
    db.add(job)
    db.flush()
    analysis = JobAnalysis(job_id=job.id, analyzer_version=ANALYZER_VERSION)
    analysis.contract_type = ContractType.WERKSTUDENT
    analysis.english_only = english_only
    analysis.raw_json = {"must_haves": ["Python"], "weekly_hours": 20}
    db.add(analysis)
    db.commit()
    return job.id


# ---- rendering -------------------------------------------------------
def test_render_produces_readable_docx(tmp_path):
    out = render_cover_letter(
        CoverLetterResult.model_validate(LETTER),
        {
            "name": "Yazan R.",
            "street": "Teststr. 1",
            "postal_code": "80331",
            "city": "München",
            "email": "y@example.com",
            "phone": "+49 1",
        },
        tmp_path / "letter.docx",
    )
    assert out.exists()
    from docx import Document

    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "Bewerbung als Werkstudent" in text
    assert "Acme GmbH" in text
    assert "ETL-Projekte" in text
    assert "München" in text


# ---- language resolution --------------------------------------
def test_resolve_language():
    de = JobAnalysis(job_id=1, analyzer_version="t", english_only=False)
    en = JobAnalysis(job_id=1, analyzer_version="t", english_only=True)
    assert _resolve_language(CoverLetterLanguageMode.ALWAYS_EN, de) == "en"
    assert _resolve_language(CoverLetterLanguageMode.MATCH_POSTING, en) == "en"
    assert _resolve_language(CoverLetterLanguageMode.MATCH_POSTING, de) == "de"
    assert _resolve_language(CoverLetterLanguageMode.MATCH_POSTING, None) == "de"


# ---- service ---------------------------------------------------
def test_generate_persists_row_docx_and_application(db_session: Session):
    job_id = _job_with_analysis(db_session)
    cl = generate(db_session, job_id, client=_client())
    db_session.commit()

    assert cl.version == 1
    assert cl.claims_used[0]["claim"] == "3 years Python"
    assert cl.docx_path and cl.docx_path.endswith(".docx")
    from pathlib import Path

    assert Path(cl.docx_path).exists()

    app = db_session.exec(select(Application).where(Application.job_id == job_id)).one()
    assert app.status == ApplicationStatus.PREPARING


def test_update_marks_edited_and_rerenders(db_session: Session):
    job_id = _job_with_analysis(db_session)
    cl = generate(db_session, job_id, client=_client())
    db_session.commit()

    new_content = dict(LETTER, subject="Bewerbung - ueberarbeitet")
    updated = update(db_session, cl.id, new_content)
    db_session.commit()
    assert updated.user_edited is True

    from docx import Document

    text = "\n".join(p.text for p in Document(updated.docx_path).paragraphs)
    assert "ueberarbeitet" in text


def test_second_generation_bumps_version(db_session: Session):
    job_id = _job_with_analysis(db_session)
    generate(db_session, job_id, client=_client())
    db_session.commit()
    second = generate(db_session, job_id, client=_client())
    db_session.commit()
    assert second.version == 2
    assert len(db_session.exec(select(CoverLetter)).all()) == 2


# ---- API -----------------------------------------------------
def test_cover_letter_api_flow(client, db_session, monkeypatch):
    job_id = _job_with_analysis(db_session)
    update_app_settings(db_session, {"cover_letter_language_mode": "always_de"})
    db_session.commit()

    monkeypatch.setattr(
        "findmyjob.services.cover_letters.generate_cover_letter",
        lambda **kw: (
            CoverLetterResult.model_validate(LETTER),
            type("R", (), {"cache_hit": False})(),
        ),
    )

    created = client.post(f"/api/jobs/{job_id}/cover-letter")
    assert created.status_code == 201, created.text
    cl_id = created.json()["id"]
    assert created.json()["claims_used"]

    listed = client.get(f"/api/jobs/{job_id}/cover-letters").json()
    assert len(listed) == 1

    edited = client.put(
        f"/api/cover-letters/{cl_id}",
        json={"content": dict(LETTER, subject="Edited subject")},
    )
    assert edited.status_code == 200 and edited.json()["user_edited"] is True

    dl = client.get(f"/api/cover-letters/{cl_id}/docx")
    assert dl.status_code == 200
    assert "wordprocessingml" in dl.headers["content-type"]
